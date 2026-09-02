from pathlib import Path

import openpyxl

from rtsprice.render import COLUMN_COUNT, C_ID, C_NAME
from rtsprice.writer import MAX_ROWS_PER_FILE, SHEET_NAME, write_price_file


def _row(rts_id: int) -> list[object]:
    row: list[object] = [None] * COLUMN_COUNT
    row[C_ID] = rts_id
    row[C_NAME] = f"Товар {rts_id}"
    return row


def test_writes_header_and_data(tmp_path: Path):
    out = tmp_path / "ooo.xlsx"
    written = write_price_file([_row(1), _row(2)], out)
    assert written == [out]

    ws = openpyxl.load_workbook(out).active
    assert ws.title == SHEET_NAME
    assert ws.cell(1, 4).value.startswith("Наименование позиции")
    assert ws.cell(2, 10).value == "Доставка транспортной компанией"
    assert ws.cell(3, 2).value == 1
    assert ws.cell(4, 2).value == 2
    assert ws.max_row == 4


def test_fills_sequential_order_numbers(tmp_path: Path):
    out = tmp_path / "ooo.xlsx"
    write_price_file([_row(10), _row(20)], out)
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(3, 1).value == 1
    assert ws.cell(4, 1).value == 2


def test_merges_delivery_header(tmp_path: Path):
    out = tmp_path / "ooo.xlsx"
    write_price_file([_row(1)], out)
    ranges = {str(r) for r in openpyxl.load_workbook(out).active.merged_cells.ranges}
    assert "J1:L1" in ranges
    assert "A1:A2" in ranges


def test_splits_into_parts_over_limit(tmp_path: Path, monkeypatch):
    monkeypatch.setattr("rtsprice.writer.MAX_ROWS_PER_FILE", 2)
    out = tmp_path / "ooo.xlsx"
    written = write_price_file([_row(i) for i in range(5)], out)
    assert [p.name for p in written] == ["ooo_part1.xlsx", "ooo_part2.xlsx", "ooo_part3.xlsx"]
    assert openpyxl.load_workbook(written[0]).active.max_row == 4
    assert openpyxl.load_workbook(written[2]).active.max_row == 3


def test_empty_rows_still_writes_header(tmp_path: Path):
    out = tmp_path / "ooo.xlsx"
    written = write_price_file([], out)
    assert openpyxl.load_workbook(written[0]).active.max_row == 2
