# tests/test_pipeline.py
from pathlib import Path

import openpyxl
import pytest

import dataclasses

from rtsprice.config import CompanyConfig, SourceConfig
from rtsprice.pipeline import Paths, build
from rtsprice.render import C_DELETE, C_ID, C_NAME, C_PRICE
from rtsprice.state import load_snapshot


def _paths(tmp_path: Path) -> Paths:
    for name in ("input", "photos", "reference", "state", "output"):
        (tmp_path / name).mkdir(exist_ok=True)
    okei = tmp_path / "reference" / "okei.csv"
    okei.write_text("code,symbol,name\n796,ШТ,Штука\n", encoding="utf-8-sig")
    (tmp_path / "reference" / "okpd2.csv").write_text("code,name\n", encoding="utf-8-sig")
    return Paths(root=tmp_path)


def _source(tmp_path: Path, rows: list[list], code: str = "s", prefix: int = 7,
            state: str = "on") -> SourceConfig:
    directory = tmp_path / "input" / code
    directory.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(["Артикул", "Наименование", "Цена"])
    for row in rows:
        ws.append(row)
    wb.save(directory / "price.xlsx")
    return SourceConfig(
        code=code, title=code.upper(), state=state, prefix=prefix, companies=("ooo_tlt",),
        file_glob=str(directory / "*.xlsx"), sheets=("Лист1",), header_rows=(1,),
        data_starts_at=2, row_is_product="price_not_empty",
        columns={"article": "Артикул", "name": "Наименование", "price": "Цена"},
    )


def _company() -> dict[str, CompanyConfig]:
    return {"ooo_tlt": CompanyConfig(code="ooo_tlt", title="ООО", vat_mode="vat22",
                                     rounding="ruble", validity_days=21)}


def test_build_writes_file_and_snapshot(tmp_path: Path):
    paths = _paths(tmp_path)
    sources = {"s": _source(tmp_path, [["A-1", "Товар", 122], ["A-2", "Второй", 244]])}
    result = build(paths, sources, _company())

    written = result.files["ooo_tlt"]
    ws = openpyxl.load_workbook(written[0]).active
    assert ws.max_row == 4
    assert ws.cell(3, C_NAME + 1).value == "Товар"
    assert (paths.state / "last_ooo_tlt.csv").exists()
    assert result.stats[0].accepted == 2


def test_second_build_marks_disappeared_position_for_deletion(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122], ["A-2", "Второй", 244]])},
          _company())
    result = build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())

    ws = openpyxl.load_workbook(result.files["ooo_tlt"][0]).active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=3)]
    deleted = [r for r in rows if r[C_DELETE] == 1]
    assert len(deleted) == 1
    assert deleted[0][C_NAME] == "Второй"
    assert result.diffs["ooo_tlt"].deleted == 1


def test_frozen_source_is_not_deleted(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())
    frozen = {"s": _source(tmp_path, [["A-1", "Товар", 122]], state="frozen")}
    result = build(paths, frozen, _company())

    ws = openpyxl.load_workbook(result.files["ooo_tlt"][0]).active
    assert ws.max_row == 2
    assert result.diffs["ooo_tlt"].deleted == 0


def test_frozen_source_can_still_be_deleted_after_being_switched_off(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]], state="frozen")},
          _company())
    result = build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]], state="off")},
                   _company())
    assert result.diffs["ooo_tlt"].deleted == 1


def test_off_source_positions_are_deleted(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())
    result = build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]], state="off")},
                   _company())
    assert result.diffs["ooo_tlt"].deleted == 1


def test_only_filter_limits_sources(tmp_path: Path):
    paths = _paths(tmp_path)
    sources = {
        "s": _source(tmp_path, [["A-1", "Товар", 122]], code="s", prefix=7),
        "t": _source(tmp_path, [["B-1", "Другой", 244]], code="t", prefix=8),
    }
    result = build(paths, sources, _company(), only=["s"])
    assert [s.code for s in result.stats] == ["s"]


def test_only_filter_does_not_delete_other_sources(tmp_path: Path):
    paths = _paths(tmp_path)
    sources = {
        "s": _source(tmp_path, [["A-1", "Товар", 122]], code="s", prefix=7),
        "t": _source(tmp_path, [["B-1", "Другой", 244]], code="t", prefix=8),
    }
    build(paths, sources, _company())
    result = build(paths, sources, _company(), only=["s"])

    assert result.diffs["ooo_tlt"].deleted == 0
    assert set(load_snapshot(paths.state / "last_ooo_tlt.csv")) == {70_000_001, 80_000_001}


def test_broken_columns_do_not_abort_build_or_delete(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())

    broken = dataclasses.replace(
        _source(tmp_path, [["A-1", "Товар", 122]]),
        columns={"article": "Артикул", "name": "Наименование", "price": "Нет такой колонки"},
    )
    result = build(paths, {"s": broken}, _company())

    assert result.stats[0].read == 0
    assert "Нет такой колонки" in " ".join(result.stats[0].reasons)
    assert result.diffs["ooo_tlt"].deleted == 0


def test_dry_run_writes_nothing(tmp_path: Path):
    paths = _paths(tmp_path)
    sources = {"s": _source(tmp_path, [["A-1", "Товар", 122]])}
    result = build(paths, sources, _company(), dry_run=True)
    assert result.files == {}
    assert not (paths.state / "last_ooo_tlt.csv").exists()
    assert result.stats[0].accepted == 1


def test_price_change_counted_in_diff(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())
    result = build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 244]])}, _company())
    assert result.diffs["ooo_tlt"].changed_price == 1
    assert result.diffs["ooo_tlt"].new == 0


def test_missing_source_file_is_reported_not_fatal(tmp_path: Path):
    paths = _paths(tmp_path)
    broken = SourceConfig(
        code="x", title="X", state="on", prefix=9, companies=("ooo_tlt",),
        file_glob=str(tmp_path / "input" / "x" / "*.xlsx"), sheets=(), header_rows=(1,),
        data_starts_at=2, row_is_product="price_not_empty", columns={},
    )
    result = build(paths, {"x": broken}, _company())
    assert result.stats[0].read == 0
    assert "не найден" in " ".join(result.stats[0].reasons)


def test_build_refuses_when_id_map_is_lost(tmp_path: Path):
    """Пустая карта при непустом снимке означала бы перепривязку всей витрины."""
    paths = _paths(tmp_path)
    sources = {"s": _source(tmp_path, [["A-1", "Товар", 122], ["A-2", "Второй", 244]])}
    build(paths, sources, _company())
    (paths.state / "id_map.csv").unlink()

    with pytest.raises(RuntimeError) as exc:
        build(paths, sources, _company())
    message = str(exc.value)
    assert "id_map.csv" in message
    assert "утрачен" in message or "потеряна" in message


def test_build_refuses_when_identifier_now_points_at_another_article(tmp_path: Path):
    paths = _paths(tmp_path)
    build(paths, {"s": _source(tmp_path, [["A-1", "Товар", 122]])}, _company())

    # Карта подменена: тот же идентификатор выдан другому артикулу.
    (paths.state / "id_map.csv").write_text(
        "source,article,rts_id\ns,ДРУГОЙ,70000001\n", encoding="utf-8-sig")

    with pytest.raises(RuntimeError) as exc:
        build(paths, {"s": _source(tmp_path, [["ДРУГОЙ", "Товар", 122]])}, _company())
    message = str(exc.value)
    assert "70000001" in message
    assert "A-1" in message and "ДРУГОЙ" in message


def test_build_refuses_before_writing_anything(tmp_path: Path):
    paths = _paths(tmp_path)
    sources = {"s": _source(tmp_path, [["A-1", "Товар", 122]])}
    build(paths, sources, _company())
    (paths.state / "id_map.csv").unlink()
    before = (paths.state / "last_ooo_tlt.csv").read_text(encoding="utf-8-sig")

    with pytest.raises(RuntimeError):
        build(paths, sources, _company())
    assert not (paths.state / "id_map.csv").exists()
    assert (paths.state / "last_ooo_tlt.csv").read_text(encoding="utf-8-sig") == before
