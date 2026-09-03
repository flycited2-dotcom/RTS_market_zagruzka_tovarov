from pathlib import Path

import openpyxl

from rtsprice.normalize import Rejection
from rtsprice.report import DiffStats, SourceStats, write_errors_xlsx, write_report_md
from rtsprice.validate import Issue


def test_errors_file_lists_rejections_and_fixes(tmp_path: Path):
    out = tmp_path / "errors.xlsx"
    write_errors_xlsx(
        out,
        [Rejection("promet", "A-1", 5, "price", "цена не число", "по запросу")],
        [("promet", "A-2", Issue("fix", "Штрих-код", "не 13 цифр, поле очищено", "123"))],
    )
    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[1]] == [
        "Источник", "Артикул", "Строка", "Уровень", "Поле", "Причина", "Значение",
    ]
    assert ws.cell(2, 4).value == "отклонено"
    assert ws.cell(2, 6).value == "цена не число"
    assert ws.cell(3, 4).value == "исправлено"


def test_report_contains_per_source_numbers(tmp_path: Path):
    out = tmp_path / "report.md"
    text = write_report_md(
        out,
        [SourceStats("promet", "Промет", "on", "прайс.xlsx", 13756, 500, 13256,
                     {"цена не число или не положительна": 500})],
        {"ooo_tlt": DiffStats(new=10, changed_price=25, unchanged=13221, deleted=3)},
    )
    assert "Промет" in text
    assert "13 256" in text or "13256" in text
    assert "цена не число" in text
    assert "новых: 10" in text
    assert "снимается: 3" in text
    assert out.read_text(encoding="utf-8") == text


def test_report_flags_source_that_lost_most_positions(tmp_path: Path):
    text = write_report_md(
        tmp_path / "report.md",
        [SourceStats("gur", "Гуриненко", "on", "p.xlsx", 100, 95, 5, {"нет цены": 95})],
        {},
    )
    assert "ВНИМАНИЕ" in text
