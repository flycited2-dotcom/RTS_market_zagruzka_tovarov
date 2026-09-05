"""Климатический хаб: Бриз, Daichi, Русклимат, JAC."""
import json
from pathlib import Path

import pytest

from rtsprice.climate import (
    ClimateError, ClimateImages, Product, fetch_catalog, save_photo_map, write_price,
)


ANSWER = {
    "ok": True,
    "total": 2,
    "products": [
        {"sku": "НС-1300165", "name": "ECOSTAR Мобильный кондиционер DESIRE KV-DS05CH-E",
         "purchasePriceGross": 15290.0, "stockQuantity": 1, "vendor": "ECOSTAR",
         "supplierName": "Бриз", "category": "7",
         "imageUrl": "https://images.breez.ru/catalog/ecostar/01.png",
         "specifications": {"series": "DESIRE"}},
        {"sku": "MDV-2", "name": "MDV Кассетный кондиционер", "purchasePriceGross": 42000.0,
         "stockQuantity": 3, "vendor": "MDV", "supplierName": "JAC", "category": "8",
         "imageUrl": "", "specifications": {"series": "AURORA"}},
    ],
}


class Server:
    host, user, key_path, port = "server", "root", "/key", 2222


def runner(answer=ANSWER):
    def run(command: str, timeout: int) -> bytes:
        assert "TENDER_CLIMATE_API_TOKEN" in command, "токен обязан читаться на сервере"
        assert "127.0.0.1:8780" in command, "сервис слушает только localhost"
        return json.dumps(answer).encode("utf-8")
    return run


def test_catalog_is_read_through_the_server():
    """Ключ хаба живёт на сервере и на рабочую машину не приезжает."""
    items = fetch_catalog(Server(), runner=runner())
    assert [i.sku for i in items] == ["НС-1300165", "MDV-2"]
    assert items[0].price == 15290.0
    assert items[0].supplier == "Бриз"
    assert items[0].series == "DESIRE"


def test_position_without_price_is_dropped():
    """Без цены строка прайса бессмысленна: показывать покупателю нечего."""
    answer = {"ok": True, "products": [
        {"sku": "X", "name": "Без цены", "purchasePriceGross": 0, "stockQuantity": 1},
    ]}
    assert fetch_catalog(Server(), runner=runner(answer)) == []


def test_position_without_article_is_dropped():
    """Артикул — ключ к номеру на площадке. Нет ключа — нет позиции."""
    answer = {"ok": True, "products": [
        {"sku": "", "name": "Без артикула", "purchasePriceGross": 100, "stockQuantity": 1},
    ]}
    assert fetch_catalog(Server(), runner=runner(answer)) == []


def test_refusal_from_the_hub_is_not_silent():
    """Отказ хаба нельзя принять за пустой каталог: иначе снимется с продажи
    всё, что он не отдал."""
    with pytest.raises(ClimateError):
        fetch_catalog(Server(), runner=runner({"ok": False, "error": "нет доступа"}))


def test_broken_answer_is_not_silent():
    def run(command: str, timeout: int) -> bytes:
        return b"<html>502</html>"
    with pytest.raises(ClimateError):
        fetch_catalog(Server(), runner=run)


def test_price_book_is_written_for_the_common_pipeline(tmp_path: Path):
    """Каталог сохраняется книгой, чтобы климат шёл теми же правилами,
    что и прайсы, присланные файлом."""
    import openpyxl

    items = fetch_catalog(Server(), runner=runner())
    path = tmp_path / "climate.xlsx"
    assert write_price(path, items) == 2
    sheet = openpyxl.load_workbook(path).active
    assert [c.value for c in sheet[1]][:4] == ["Артикул", "Наименование", "Цена", "Остаток"]
    assert sheet.cell(row=2, column=1).value == "НС-1300165"
    assert sheet.cell(row=2, column=3).value == 15290.0


def test_photo_map_keeps_only_positions_with_a_picture(tmp_path: Path):
    items = fetch_catalog(Server(), runner=runner())
    path = tmp_path / "photos.json"
    assert save_photo_map(path, items) == 1
    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved == {"НС-1300165": "https://images.breez.ru/catalog/ecostar/01.png"}


def test_images_index_answers_by_article():
    index = ClimateImages({"НС-1": "https://images.breez.ru/a.png", "НС-2": ""})
    assert index.images(["НС-1", "НС-2", "НС-3"]) == {
        "НС-1": "https://images.breez.ru/a.png"}


def test_missing_photo_map_explains_what_to_do(tmp_path: Path):
    with pytest.raises(ClimateError, match="rtsprice climate"):
        ClimateImages.from_file(tmp_path / "нет.json")
