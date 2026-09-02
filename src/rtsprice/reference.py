"""Справочники площадки: единицы измерения ОКЕИ и коды ОКПД2."""
from __future__ import annotations

import csv
import re
from pathlib import Path

import yaml

UNITS_SHEET = "Справочник единиц измерения"
OKPD2_SHEET = "Справочник ОКПД2"
OKPD2_PATTERN = re.compile(r"^\d{2}(\.\d{1,2}){0,2}(\.\d{1,3})?$")


def _key(value: str) -> str:
    """Ключ сравнения: без регистра, без лишних пробелов и точек по краям."""
    return re.sub(r"\s+", " ", str(value)).strip().strip(".").casefold()


def extract_reference_tables(template: Path, out_dir: Path) -> None:
    """Единожды вытащить справочники из шаблона РТС в CSV."""
    import openpyxl

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(template, data_only=True, read_only=True)

    with (out_dir / "okei.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "symbol", "name"])
        for symbol, name, code in wb[UNITS_SHEET].iter_rows(min_row=2, max_col=3, values_only=True):
            if symbol and str(symbol).strip() and str(symbol).strip() != "-":
                w.writerow([code, str(symbol).strip(), str(name or "").strip()])

    with (out_dir / "okpd2.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "name"])
        for code, name in wb[OKPD2_SHEET].iter_rows(min_row=2, max_col=2, values_only=True):
            if code:
                w.writerow([str(code).strip(), str(name or "").strip()])
    wb.close()


def load_units(path: Path) -> set[str]:
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return {row["symbol"].strip() for row in csv.DictReader(fh) if row.get("symbol")}


def load_unit_aliases(path: Path) -> dict[str, str]:
    p = Path(path)
    if not p.exists():
        return {}
    raw = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    return {_key(k): str(v) for k, v in raw.items()}


def resolve_unit(raw: object, units: set[str], aliases: dict[str, str]) -> str | None:
    """Привести произвольное обозначение к обозначению ОКЕИ или вернуть None."""
    if raw is None or not str(raw).strip():
        return None
    index = {_key(u): u for u in units}
    k = _key(raw)
    if k in index:
        return index[k]
    if k in aliases:
        return aliases[k]
    return None


def load_okpd2(path: Path) -> set[str]:
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return {row["code"].strip() for row in csv.DictReader(fh) if row.get("code")}
