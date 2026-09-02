"""Запись выходного файла в формате шаблона РТС."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .render import (
    COLUMN_COUNT, C_DELIVERY_CARRIER, C_ORDER, C_PICKUP, HEADER_ROW_1, HEADER_ROW_2,
)

MAX_ROWS_PER_FILE = 50_000
SHEET_NAME = "Данные для импорта"


def _new_book() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for column, value in enumerate(HEADER_ROW_1, start=1):
        ws.cell(1, column, value)
    for column, value in enumerate(HEADER_ROW_2, start=1):
        ws.cell(2, column, value)

    first = get_column_letter(C_DELIVERY_CARRIER + 1)
    last = get_column_letter(C_PICKUP + 1)
    ws.merge_cells(f"{first}1:{last}1")
    for index in range(COLUMN_COUNT):
        if C_DELIVERY_CARRIER <= index <= C_PICKUP:
            continue
        letter = get_column_letter(index + 1)
        ws.merge_cells(f"{letter}1:{letter}2")
    return wb


def _write_chunk(rows: list[list[object]], path: Path, start_order: int) -> Path:
    wb = _new_book()
    ws = wb.active
    for offset, row in enumerate(rows):
        values = list(row)
        values[C_ORDER] = start_order + offset
        ws.append(values)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_price_file(rows: list[list[object]], out_path: Path) -> list[Path]:
    """Записать строки, разбив на части, если их больше допустимого предела."""
    out_path = Path(out_path)
    if len(rows) <= MAX_ROWS_PER_FILE:
        return [_write_chunk(rows, out_path, 1)]

    written: list[Path] = []
    for part, start in enumerate(range(0, len(rows), MAX_ROWS_PER_FILE), start=1):
        chunk = rows[start:start + MAX_ROWS_PER_FILE]
        path = out_path.with_name(f"{out_path.stem}_part{part}{out_path.suffix}")
        written.append(_write_chunk(chunk, path, start + 1))
    return written
